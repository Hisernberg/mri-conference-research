#!/usr/bin/env python3
"""Convert project Markdown to styled DOCX files for an installed file exporter.

The Adobe text-to-PDF service returned internal errors in this session. The
documents are rendered locally by LibreOffice from these existing DOCX files.
This script writes DOCX, not PDF, and does not contact any external service.
"""
from pathlib import Path
import os
import re
from urllib.parse import quote, unquote, urlsplit, urlunsplit
import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, element, source_path):
    for node in element.children:
        if isinstance(node, NavigableString):
            paragraph.add_run(str(node))
        elif node.name == "a":
            # Preserve external URLs; rebase local targets for docs/docx and docs/pdf.
            href = node.get("href", "")
            parts = urlsplit(href)
            if not parts.scheme and not parts.netloc and parts.path:
                target = (source_path.parent / unquote(parts.path)).resolve()
                relative = os.path.relpath(target, ROOT / "docs" / "docx")
                href = urlunsplit(("", "", quote(relative, safe="/"), parts.query, parts.fragment))
            part = paragraph.part
            rel = part.relate_to(href,
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rel)
            run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "176B87"); props.append(color)
            run.append(props); text = OxmlElement("w:t"); text.text = node.get_text(); run.append(text)
            link.append(run); paragraph._p.append(link)
        else:
            run = paragraph.add_run(node.get_text())
            run.bold = node.name in {"strong", "b"}
            run.italic = node.name in {"em", "i"}
            if node.name == "code":
                run.font.name = "Liberation Mono"; run.font.size = Pt(8)


def convert(path):
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65); section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.25); section.footer_distance = Inches(0.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for level, size in [(1, 21), (2, 14), (3, 11)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Liberation Sans"; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("153E52")
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "MRI RESEARCH  |  EXPERIMENT RECORD"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.add_run("2026-09-09  |  ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    html = markdown.markdown(path.read_text(), extensions=["tables", "fenced_code"])
    soup = BeautifulSoup(html, "html.parser")
    for node in soup.children:
        if isinstance(node, NavigableString):
            continue
        if re.fullmatch(r"h[1-6]", node.name or ""):
            if path.stem == "FIGURE_CATALOG" and node.name == "h2" and node.get_text().startswith("Figure "):
                doc.add_page_break()
            doc.add_heading(node.get_text(), min(int(node.name[1]), 3))
        elif node.name == "p":
            img = node.find("img")
            if img is not None:
                image_path = (path.parent / img.get("src", "")).resolve()
                if image_path.is_file():
                    doc.add_picture(str(image_path), width=Inches(6.65))
            else:
                inline(doc.add_paragraph(), node, path)
        elif node.name in {"ul", "ol"}:
            for li in node.find_all("li", recursive=False):
                inline(doc.add_paragraph(style="List Bullet" if node.name == "ul" else "List Number"), li, path)
        elif node.name == "pre":
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
            run = p.add_run(node.get_text().rstrip()); run.font.name = "Liberation Mono"; run.font.size = Pt(7.5)
        elif node.name == "table":
            rows = node.find_all("tr"); columns = len(rows[0].find_all(["th", "td"]))
            table = doc.add_table(rows=0, cols=columns); table.style = "Light Shading Accent 1"
            headers = [cell.get_text().strip() for cell in rows[0].find_all(["th", "td"])]
            keep_table = len(rows) <= 6 or headers[0] == "Available modalities"
            if headers == ["Job", "Version", "Evidence"]:
                table.autofit = False
                available = section.page_width - section.left_margin - section.right_margin
                for column, fraction in zip(table.columns, [.27, .10, .63]):
                    column.width = int(available * fraction)
            for i, tr in enumerate(rows):
                cells = table.add_row().cells
                no_split = OxmlElement("w:cantSplit")
                table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                for cell, item in zip(cells, tr.find_all(["th", "td"], recursive=False)):
                    inline(cell.paragraphs[0], item, path)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.keep_with_next = i == 0 or (keep_table and i < len(rows) - 1)
                        p.paragraph_format.keep_together = True
                        for run in p.runs:
                            run.font.size = Pt(8.5); run.bold = bool(i == 0) or run.bold
                if i == 0:
                    repeat = OxmlElement("w:tblHeader"); table.rows[i]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
    out = ROOT / "docs" / "docx" / (path.stem + ".docx"); out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out); print(out)


if __name__ == "__main__":
    for path in [ROOT / "EXPERIMENT_PLAN.md", *sorted((ROOT / "docs").glob("*.md"))]:
        convert(path)
